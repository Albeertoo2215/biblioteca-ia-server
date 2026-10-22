"""Utilidades comunes para las hojas de cálculo de Casitres (openpyxl)."""
import os
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
AQUI = os.path.dirname(os.path.abspath(__file__))
MARINO, AZUL, GRIS, LINEA = '0F2340', '2F6FDE', 'F3F6FA', 'E1E7EF'
F_TIT = Font(name='Arial', size=20, bold=True, color=MARINO)
F_SUB = Font(name='Arial', size=11, color='5B6B82')
F_HEAD = Font(name='Arial', size=10, bold=True, color='FFFFFF')
F_TXT = Font(name='Arial', size=10, color=MARINO)
F_BOLD = Font(name='Arial', size=10, bold=True, color=MARINO)
F_PIE = Font(name='Arial', size=8, italic=True, color='5B6B82')
FILL_HEAD = PatternFill('solid', fgColor=AZUL)
FILL_MAR = PatternFill('solid', fgColor=MARINO)
FILL_GRIS = PatternFill('solid', fgColor=GRIS)
FILL_INPUT = PatternFill('solid', fgColor='FFF8E6')
BORDE = Border(bottom=Side(style='thin', color=LINEA))
EUR = '#,##0.00 €'

def cabecera(ws, titulo, subtitulo, ancho_cols):
    """Logo casi3, título, subtítulo y aviso «Creado con IA» arriba de cada hoja."""
    ws.sheet_view.showGridLines = False
    for col, w in ancho_cols.items():
        ws.column_dimensions[col].width = w
    img = XLImage(os.path.join(AQUI, 'logo.png')); img.height = 26; img.width = int(26 * 577 / 200)
    ws.add_image(img, 'B1'); ws.row_dimensions[1].height = 24
    ws['B3'] = titulo; ws['B3'].font = F_TIT; ws.row_dimensions[3].height = 28
    ws['B4'] = subtitulo; ws['B4'].font = F_SUB
    ws['B5'] = 'Creado con IA · casi3 · Las celdas amarillas son las que rellenas tú.'; ws['B5'].font = F_PIE

def encabezados(ws, fila, col0, textos, fill=FILL_HEAD):
    for i, t in enumerate(textos):
        c = ws.cell(row=fila, column=col0 + i, value=t)
        c.font = F_HEAD; c.fill = fill; c.alignment = Alignment(vertical='center', wrap_text=True)
    ws.row_dimensions[fila].height = 22

def celda(ws, ref, valor=None, font=F_TXT, fmt=None, entrada=False, borde=True):
    c = ws[ref]
    if valor is not None: c.value = valor
    c.font = font
    if fmt: c.number_format = fmt
    if entrada: c.fill = FILL_INPUT
    if borde: c.border = BORDE
    return c
