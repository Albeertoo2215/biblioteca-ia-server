"""Presupuesto mensual: hoja de cálculo (Excel / Google Sheets) + guía en PDF."""
import os, sys
sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', 'comun'))
from openpyxl import Workbook
from openpyxl.chart import BarChart, PieChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, DataBarRule, FormulaRule
from openpyxl.styles import PatternFill, Font
from excel import *
from plantilla import documento, donut, barras
from construir import pdf

DIST = os.path.join(os.path.dirname(__file__), '..', '..', 'dist')
wb = Workbook()

# ── Empieza aquí ──
ws = wb.active; ws.title = 'Empieza aquí'
cabecera(ws, 'Presupuesto mensual', 'Controla en qué se va tu dinero en 10 minutos al mes.', {'A': 3, 'B': 30, 'C': 70})
pasos = [('1. Ingresos', 'En la hoja «Mes», escribe lo que cobras (nómina, extras). Solo las celdas amarillas.'),
         ('2. Presupuesto', 'Pon cuánto quieres gastar en cada categoría. Elige si es Necesidad, Deseo o Ahorro.'),
         ('3. Gastos reales', 'Durante el mes, anota lo que gastas de verdad en la columna «Real».'),
         ('4. Revisa', 'La diferencia sale en rojo si te pasas. El bloque 50/30/20 te dice si vas equilibrado.'),
         ('5. Año', 'A final de mes, copia los totales en la hoja «Año» y verás tu evolución en el gráfico.'),
         ('6. Objetivos y deudas', 'Apunta tus metas de ahorro y tus deudas: la hoja calcula cuánto apartar cada mes.'),
         ('En Google Sheets', 'Sube el archivo a Google Drive y ábrelo con Hojas de cálculo de Google. Todo funciona igual.')]
for i, (a, b) in enumerate(pasos):
    r = 8 + i * 2
    celda(ws, f'B{r}', a, F_BOLD); celda(ws, f'C{r}', b); ws.row_dimensions[r].height = 30
    ws[f'C{r}'].alignment = Alignment(wrap_text=True, vertical='top')

# ── Mes ──
m = wb.create_sheet('Mes')
cabecera(m, 'Mi mes', 'Ingresos, presupuesto y gasto real por categoría.', {'A': 3, 'B': 26, 'C': 15, 'D': 15, 'E': 15, 'F': 14, 'G': 3, 'H': 22, 'I': 15, 'J': 15, 'K': 12})
encabezados(m, 7, 2, ['Ingresos', 'Importe'])
ingresos = [('Nómina', 1650), ('Extras', 0), ('Otros ingresos', 0)]
for i, (k, v) in enumerate(ingresos):
    r = 8 + i; celda(m, f'B{r}', k, entrada=True); celda(m, f'C{r}', v, fmt=EUR, entrada=True)
celda(m, 'B11', 'Total ingresos', F_BOLD); celda(m, 'C11', '=SUM(C8:C10)', F_BOLD, EUR)

encabezados(m, 13, 2, ['Categoría', 'Tipo', 'Presupuesto', 'Real', 'Diferencia'])
cats = [('Alquiler o hipoteca', 'Necesidad', 650), ('Luz, agua y gas', 'Necesidad', 90), ('Internet y móvil', 'Necesidad', 45),
        ('Supermercado', 'Necesidad', 260), ('Transporte', 'Necesidad', 60), ('Seguros', 'Necesidad', 30), ('Salud y farmacia', 'Necesidad', 20),
        ('Restaurantes y bares', 'Deseo', 80), ('Ocio y planes', 'Deseo', 50), ('Suscripciones', 'Deseo', 25), ('Ropa', 'Deseo', 40),
        ('Regalos y otros', 'Deseo', 20), ('Ahorro', 'Ahorro', 200), ('Fondo de emergencia', 'Ahorro', 50), ('Pago extra de deudas', 'Ahorro', 30)]
dv = DataValidation(type='list', formula1='"Necesidad,Deseo,Ahorro"', allow_blank=True); m.add_data_validation(dv)
r0 = 14
for i, (k, t, p) in enumerate(cats):
    r = r0 + i
    celda(m, f'B{r}', k, entrada=True); celda(m, f'C{r}', t, entrada=True); dv.add(f'C{r}')
    celda(m, f'D{r}', p, fmt=EUR, entrada=True); celda(m, f'E{r}', None, fmt=EUR, entrada=True)
    celda(m, f'F{r}', f'=D{r}-E{r}', fmt=EUR)
r1 = r0 + len(cats) - 1; rt = r1 + 1
celda(m, f'B{rt}', 'Total', F_BOLD)
for col in 'DEF': celda(m, f'{col}{rt}', f'=SUM({col}{r0}:{col}{r1})', F_BOLD, EUR)
m.conditional_formatting.add(f'F{r0}:F{rt}', CellIsRule(operator='lessThan', formula=['0'], font=Font(color='C0261B', bold=True), fill=PatternFill('solid', fgColor='FDECEA')))
m.conditional_formatting.add(f'E{r0}:E{r1}', DataBarRule(start_type='num', start_value=0, end_type='max', color='2F6FDE'))
celda(m, f'B{rt+2}', 'Te queda este mes', F_BOLD); celda(m, f'C{rt+2}', f'=C11-E{rt}', Font(name='Arial', size=12, bold=True, color='1F8A7A'), EUR)

# Bloque 50/30/20
encabezados(m, 7, 8, ['Regla 50/30/20', 'Recomendado', 'Tu reparto', '% real'], FILL_MAR)
for i, (t, pct) in enumerate([('Necesidad', .5), ('Deseo', .3), ('Ahorro', .2)]):
    r = 8 + i
    celda(m, f'H{r}', t, F_BOLD); celda(m, f'I{r}', f'=$C$11*{pct}', fmt=EUR)
    celda(m, f'J{r}', f'=SUMIF($C${r0}:$C${r1},H{r},$E${r0}:$E${r1})', fmt=EUR)
    celda(m, f'K{r}', f'=IF($C$11=0,0,J{r}/$C$11)', fmt='0%')
m['H12'] = 'Si «Necesidad» pasa del 50 %, busca dónde recortar (tarifas, seguros, suscripciones).'; m['H12'].font = F_PIE

ch = BarChart(); ch.type = 'bar'; ch.style = 10; ch.title = 'Presupuesto frente a gasto real'; ch.height = 9; ch.width = 17
ch.add_data(Reference(m, min_col=4, max_col=5, min_row=13, max_row=r1), titles_from_data=True)
ch.set_categories(Reference(m, min_col=2, min_row=r0, max_row=r1)); ch.y_axis.numFmt = '0 €'
ch.series[0].graphicalProperties.solidFill = 'C9D8F5'; ch.series[1].graphicalProperties.solidFill = AZUL
m.add_chart(ch, 'H15')
pie = PieChart(); pie.title = 'Reparto del gasto real'; pie.height = 8; pie.width = 12
pie.add_data(Reference(m, min_col=10, min_row=7, max_row=10), titles_from_data=True); pie.set_categories(Reference(m, min_col=8, min_row=8, max_row=10))
pie.dataLabels = DataLabelList(); pie.dataLabels.showPercent = True
m.add_chart(pie, 'H34')

# ── Año ──
a = wb.create_sheet('Año')
cabecera(a, 'Mi año', 'Copia aquí los totales de cada mes y mira la tendencia.', {'A': 3, 'B': 16, 'C': 15, 'D': 15, 'E': 15, 'F': 15})
encabezados(a, 7, 2, ['Mes', 'Ingresos', 'Gastos', 'Ahorrado', 'Ahorro acumulado'])
meses = ['Enero','Febrero','Marzo','Abril','Mayo','Junio','Julio','Agosto','Septiembre','Octubre','Noviembre','Diciembre']
for i, mm in enumerate(meses):
    r = 8 + i
    celda(a, f'B{r}', mm, F_BOLD); celda(a, f'C{r}', None, fmt=EUR, entrada=True); celda(a, f'D{r}', None, fmt=EUR, entrada=True)
    celda(a, f'E{r}', f'=IF(C{r}="","",C{r}-D{r})', fmt=EUR); celda(a, f'F{r}', f'=IF(E{r}="","",SUM($E$8:E{r}))', fmt=EUR)
celda(a, 'B20', 'Total', F_BOLD)
for col in 'CDE': celda(a, f'{col}20', f'=SUM({col}8:{col}19)', F_BOLD, EUR)
lc = LineChart(); lc.title = 'Ingresos y gastos'; lc.height = 8; lc.width = 18; lc.y_axis.numFmt = '0 €'
lc.add_data(Reference(a, min_col=3, max_col=4, min_row=7, max_row=19), titles_from_data=True); lc.set_categories(Reference(a, min_col=2, min_row=8, max_row=19))
a.add_chart(lc, 'H7')

# ── Objetivos ──
o = wb.create_sheet('Objetivos de ahorro')
cabecera(o, 'Objetivos de ahorro', 'Cuánto apartar cada mes para llegar a tiempo.', {'A': 3, 'B': 26, 'C': 14, 'D': 14, 'E': 14, 'F': 13, 'G': 16})
encabezados(o, 7, 2, ['Objetivo', 'Necesito', 'Para (fecha)', 'Llevo', 'Progreso', 'Apartar al mes'])
import datetime as dt
ejemplos = [('Fondo de emergencia (3 meses de gastos)', 4200, dt.date(2027, 12, 31), 600), ('Viaje de verano', 1200, dt.date(2027, 7, 1), 150), ('Ordenador nuevo', 900, dt.date(2027, 4, 1), 0)]
for i in range(8):
    r = 8 + i; e = ejemplos[i] if i < len(ejemplos) else (None, None, None, None)
    celda(o, f'B{r}', e[0], entrada=True); celda(o, f'C{r}', e[1], fmt=EUR, entrada=True); celda(o, f'D{r}', e[2], fmt='dd/mm/yyyy', entrada=True)
    celda(o, f'E{r}', e[3], fmt=EUR, entrada=True)
    celda(o, f'F{r}', f'=IF(C{r}="","",MIN(1,E{r}/C{r}))', fmt='0%')
    celda(o, f'G{r}', f'=IF(OR(C{r}="",D{r}=""),"",MAX(0,(C{r}-E{r})/MAX(1,DATEDIF(TODAY(),D{r},"m"))))', fmt=EUR)
o.conditional_formatting.add('F8:F15', DataBarRule(start_type='num', start_value=0, end_type='num', end_value=1, color='1F8A7A'))

# ── Deudas ──
d = wb.create_sheet('Deudas')
cabecera(d, 'Mis deudas', 'Método bola de nieve: paga el mínimo en todas y el extra a la más pequeña.', {'A': 3, 'B': 24, 'C': 14, 'D': 12, 'E': 14, 'F': 16, 'G': 12})
encabezados(d, 7, 2, ['Deuda', 'Saldo', 'Interés anual', 'Cuota mínima', 'Meses para acabar', 'Orden'])
deu = [('Tarjeta de crédito', 1200, .2, 60), ('Préstamo coche', 6500, .07, 180), ('Compra a plazos', 400, 0, 50)]
for i in range(6):
    r = 8 + i; e = deu[i] if i < len(deu) else (None, None, None, None)
    celda(d, f'B{r}', e[0], entrada=True); celda(d, f'C{r}', e[1], fmt=EUR, entrada=True); celda(d, f'D{r}', e[2], fmt='0.0%', entrada=True); celda(d, f'E{r}', e[3], fmt=EUR, entrada=True)
    celda(d, f'F{r}', f'=IF(OR(C{r}="",E{r}=""),"",IF(D{r}=0,ROUNDUP(C{r}/E{r},0),IFERROR(ROUNDUP(NPER(D{r}/12,-E{r},C{r}),0),"La cuota no cubre los intereses")))')
    celda(d, f'G{r}', f'=IF(C{r}="","",RANK(C{r},$C$8:$C$13,1))')
celda(d, 'B15', 'Total', F_BOLD); celda(d, 'C15', '=SUM(C8:C13)', F_BOLD, EUR); celda(d, 'E15', '=SUM(E8:E13)', F_BOLD, EUR)
d['B17'] = 'El orden 1 es la deuda que atacas primero con el dinero extra. Cuando la terminas, su cuota pasa a la siguiente.'; d['B17'].font = F_PIE

for s in wb.worksheets: s.sheet_properties.tabColor = AZUL if s.title != 'Empieza aquí' else MARINO
os.makedirs(DIST, exist_ok=True)
wb.save(os.path.join(DIST, 'Presupuesto mensual - casi3.xlsx')); print('XLSX ok')

# ── Guía en PDF ──
p1 = f'''<h2>Tu dinero, en una hoja</h2><p class="lead">La mayoría de la gente no gasta de más por capricho: gasta de más porque no sabe en qué se le va. Esta plantilla lo hace visible en 10 minutos al mes.</p>
<div class="grid2" style="align-items:center;margin-top:6mm"><div>{donut([('Necesidades', 50, '#0f2340'), ('Deseos', 30, '#2f6fde'), ('Ahorro', 20, '#1f8a7a')], '50/30/20')}</div>
<div><h3>La regla 50/30/20</h3><p><b style="color:#0f2340">50 % necesidades:</b> casa, suministros, comida, transporte, seguros.</p><p><b style="color:#2f6fde">30 % deseos:</b> restaurantes, ocio, ropa, suscripciones.</p><p><b style="color:#1f8a7a">20 % ahorro:</b> fondo de emergencia, objetivos y pagar deudas antes de tiempo.</p>
<div class="callout">Es un punto de partida, no una ley. Con un alquiler alto es normal que las necesidades pasen del 50 %: lo importante es saberlo y decidir dónde compensar.</div></div></div>'''
p2 = '''<h2>Cómo usar la plantilla</h2><ol class="steps">
<li><b>Abre el archivo.</b> En Excel directamente, o súbelo a Google Drive y ábrelo con Hojas de cálculo de Google.</li>
<li><b>Rellena solo lo amarillo.</b> Todo lo demás se calcula solo: diferencias, totales, porcentajes y gráficos.</li>
<li><b>Hoja «Mes».</b> Escribe tus ingresos, tu presupuesto por categoría y, durante el mes, lo que gastas de verdad.</li>
<li><b>Mira los rojos.</b> Si una diferencia sale en rojo, te has pasado en esa categoría. El bloque 50/30/20 te dice si el conjunto está equilibrado.</li>
<li><b>Hoja «Año».</b> El último día del mes copia el total de ingresos y gastos. El gráfico te enseña la tendencia.</li>
<li><b>Objetivos y deudas.</b> Indica cuánto necesitas y para cuándo: la hoja calcula cuánto apartar al mes. En deudas, te ordena cuál atacar primero.</li></ol>
<div class="callout"><b>Truco:</b> pon el «Ahorro» como un gasto más y transfiérelo el día que cobras. Lo que sobra a final de mes casi nunca llega a ahorrarse.</div>'''
p3 = f'''<h2>Dónde suele haber dinero escondido</h2><p class="lead">Revisa estas partidas una vez al año. Son las que más se pagan por inercia.</p>
<div class="grid2">
<div class="card"><h3>Luz y gas</h3><p>Compara tu tarifa con el comparador oficial de ofertas de la CNMC (comparador.cnmc.gob.es). Mira también la potencia contratada: muchas casas pagan más potencia de la que usan.</p></div>
<div class="card"><h3>Móvil e internet</h3><p>Las ofertas de permanencia caducan y la cuota sube sin avisar. Llama al terminar la permanencia o compara operadoras.</p></div>
<div class="card"><h3>Seguros</h3><p>Coche, hogar y vida se renuevan solos. Pide presupuesto a otras dos aseguradoras un mes antes de la renovación.</p></div>
<div class="card"><h3>Suscripciones</h3><p>Repasa el extracto del banco: plataformas, apps y gimnasio que ya no usas. Cancela antes de la fecha de renovación.</p></div>
<div class="card"><h3>Comisiones del banco</h3><p>Si tu banco te cobra mantenimiento, pregunta qué condiciones lo eliminan o cambia a una cuenta sin comisiones.</p></div>
<div class="card"><h3>Supermercado</h3><p>Planifica el menú de la semana y compra con lista. Comprar sin plan es la forma más rápida de tirar comida y dinero.</p></div></div>'''
p4 = '''<h2>El fondo de emergencia</h2><p class="lead">Es el primer objetivo de ahorro: dinero apartado solo para imprevistos (avería, paro, dentista).</p>
<div class="grid3" style="margin:5mm 0"><div class="card"><div class="num">3</div><p>meses de gastos fijos como mínimo.</p></div><div class="card"><div class="num">6</div><p>meses si eres autónomo o tus ingresos varían.</p></div><div class="card"><div class="num">1</div><p>cuenta separada, para no gastarlo sin querer.</p></div></div>
<h3>Cómo calcularlo</h3><p>Suma en la hoja «Mes» solo las categorías de tipo Necesidad. Multiplica por 3 (o por 6). Ese es tu objetivo: apúntalo en «Objetivos de ahorro» y la hoja te dirá cuánto apartar cada mes.</p>
<h3>Bola de nieve para deudas</h3><p>Paga la cuota mínima de todas. Todo el dinero extra va a la deuda más pequeña. Cuando la terminas, su cuota se suma a la siguiente. Avanzas rápido y ves resultados pronto, que es lo que hace que no lo dejes.</p>
<div class="callout warn"><b>Importante:</b> esta plantilla es una herramienta de organización. No es asesoramiento financiero. Para decisiones de inversión o de deuda importantes, consulta con un profesional.</div>'''
html = documento('Presupuesto mensual', 'Guía de uso y trucos para que tu dinero rinda. Acompaña a la hoja de cálculo.', 'Plantillas · casi3', 'marino', [p1, p2, p3, p4])
pdf(html, os.path.join(DIST, 'Presupuesto mensual - guia de uso - casi3.pdf'), 'Presupuesto mensual')
